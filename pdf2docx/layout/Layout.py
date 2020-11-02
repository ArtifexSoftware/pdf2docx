# -*- coding: utf-8 -*-

'''
Layout objects based on PDF raw dict extracted with PyMuPDF.

@created: 2020-07-22
@author: train8808@gmail.com
---

The raw page content extracted with PyMuPDF, `page.getText('rawdict')` is described per link:
https://pymupdf.readthedocs.io/en/latest/textpage.html

In addition to the raw layout dict, some new features are also included, e.g.
    - page margin
    - rectangle shapes, for text format, annotations and table border/shading
    - new block in table type

{
    # raw dict
    ----------------------------
    "width" : w,
    "height": h,    
    "blocks": [{...}, {...}, ...],

    # introduced dict
    ----------------------------
    "margin": [left, right, top, bottom],
    "shapes" : [{...}, {...}, ...]
}
'''


import json
from collections import defaultdict
from docx.shared import Pt
from docx.enum.section import WD_SECTION
from .Blocks import Blocks
from ..image.Image import ImagesExtractor
from ..shape.Shapes import Shapes
from ..shape.Paths import Paths
from ..table.TablesConstructor import TablesConstructor
from ..common.BBox import BBox
from ..common.share import debug_plot
from ..common import constants


class Layout:
    ''' Object representing the whole page, e.g. margins, blocks, shapes, spacing.'''

    def __init__(self, parent=None, settings:dict=None):
        ''' Initialize page layout.
            ---
            Args:
            - parent: fitz.Page, owner of this layout
            - settings: global parameters for layout parsing
        '''
        # global configuration parameters
        self.settings = self.__init_settings(settings)
       
        # initialize layout
        data = self.__source_from_page(parent) if parent else {}
        self.restore(data)

        # whether this layout is finalized already
        self.finalized = data.get('finalized', False)

        # plot initial layout for debug purpose: settings['debug']=True
        self.plot()


    @staticmethod
    def __init_settings(settings:dict):
        default = {
            'debug': False, # plot layout if True
            'connected_border_tolerance'     : 0.5, # two borders are intersected if the gap lower than this value
            'max_border_width'               : 6.0, # max border width
            'min_border_clearance'           : 2.0, # the minimum allowable clearance of two borders
            'float_image_ignorable_gap'      : 5.0, # float image if the intersection exceeds this value
            'float_layout_tolerance'         : 0.1, # [0,1] the larger of this value, the more tolerable of float layout
            'page_margin_factor_top'         : 0.5, # [0,1] reduce top margin by factor
            'page_margin_factor_bottom'      : 0.5, # [0,1] reduce bottom margin by factor
            'shape_merging_threshold'        : 0.5, # [0,1] merge shape if the intersection exceeds this value
            'shape_min_dimension'            : 2.0, # ignore shape if both width and height is lower than this value
            'line_overlap_threshold'         : 0.9, # [0,1] delete line if the intersection to other lines exceeds this value
            'line_merging_threshold'         : 2.0, # combine two lines if the x-distance is lower than this value
            'line_separate_threshold'        : 5.0, # two separate lines if the x-distance exceeds this value
            'lines_left_aligned_threshold'   : 1.0, # left aligned if delta left edge of two lines is lower than this value
            'lines_right_aligned_threshold'  : 1.0, # right aligned if delta right edge of two lines is lower than this value
            'lines_center_aligned_threshold' : 2.0, # center aligned if delta center of two lines is lower than this value
            'clip_image_res_ratio'           : 3.0, # resolution ratio (to 72dpi) when cliping page image
            'curve_path_ratio'               : 0.2, # clip page bitmap if the component of curve paths exceeds this ratio
            'extract_stream_table'           : False, # don't consider stream table when extracting tables
        }

        # update user defined parameters
        if settings: default.update(settings)
        return default


    @property
    def bbox(self): return (0.0, 0.0, self.width, self.height)


    @property
    def working_bbox(self):
        '''bbox with margin considered.'''
        x0, y0, x1, y1 = self.bbox
        L, R, T, B = self.margin
        return (x0+L, y0+T, x1-R, y1-B)


    @property
    def margin(self):
        '''Calculate page margin, (left, right, top, bottom).'''
        # return normal page margin if no blocks exist
        if not self.blocks and not self.shapes: return (constants.ITP, ) * 4

        x0, y0, x1, y1 = self.bbox
        u0, v0, u1, v1 = self.blocks.bbox | self.shapes.bbox

        # margin
        left = max(u0-x0, 0.0)
        right = max(x1-u1-constants.MINOR_DIST, 0.0)
        top = max(v0-y0, 0.0)
        bottom = max(y1-v1, 0.0)

        # reduce calculated top/bottom margin to leave some free space
        top *= self.settings['page_margin_factor_top']
        bottom *= self.settings['page_margin_factor_bottom']

        # use normal margin if calculated margin is large enough
        return (
            min(constants.ITP, left), 
            min(constants.ITP, right), 
            min(constants.ITP, top), 
            min(constants.ITP, bottom)
            )


    def reset(self):
        '''Reset Layout object.'''
        self.finalized = False

        # blocks representing text/table contents
        self.blocks = Blocks(parent=self)

        # shapes representing table border, shading and text style like underline, highlight
        self.shapes = Shapes(parent=self)

        # table parser
        self._tables_constructor = TablesConstructor(parent=self)


    def parse(self):
        ''' Parse page layout.'''
        # parse layout only once
        if self.finalized: return self

        # preprocessing, e.g. change block order, clean negative block
        self.clean_up_blocks()
        self.clean_up_shapes() # based on cleaned blocks
    
        # parse table blocks: 
        #  - table structure/format recognized from rectangles
        self.parse_lattice_tables()
        
        #  - cell contents extracted from text blocks
        self.parse_stream_tables()

        # parse text format, e.g. highlight, underline
        self.parse_text_format()
        
        # paragraph / line spacing        
        self.parse_spacing()

        self.finalized = True

        return self


    def store(self):
        '''Store parsed layout.'''
        return {
            'finalized': self.finalized,
            'width' : self.width,
            'height': self.height,
            'margin': self.margin,
            'blocks': self.blocks.store(),
            'shapes': self.shapes.store(),
        }


    def restore(self, data:dict):
        '''Restore Layout from parsed results.'''
        # reset attributes
        self.reset()

        # page width/height
        self.width = data.get('width', 0.0)
        self.height = data.get('height', 0.0)
        
        # initialize blocks
        self.blocks.restore(data.get('blocks', []))

        # initialize shapes: to add rectangles later
        self.shapes.restore(data.get('shapes', []))

        return self


    def serialize(self, filename:str):
        '''Write layout to specified file.'''
        with open(filename, 'w', encoding='utf-8') as f:
            f.write(json.dumps(self.store(), indent=4))


    def extract_tables(self):
        '''Extract content from tables.'''
        # check table
        tables = [] # type: list[ list[list[str]] ]
        if self.settings['extract_stream_table']:
            collections = self.blocks.table_blocks
        else:
            collections = self.blocks.lattice_table_blocks
        
        for table_block in collections:
            tables.append(table_block.text)

        return tables


    def make_page(self, doc):
        ''' Create page based on layout data. 
            ---
            Args:
            - doc: python-docx.Document object

            To avoid incorrect page break from original document, a new page section
            is created for each page.

            The vertical postion of paragraph/table is defined by space_before or 
            space_after property of a paragraph.
        '''
        # new page section
        # a default section is created when initialize the document,
        # so we do not have to add section for the first time.
        if not doc.paragraphs:
            section = doc.sections[0]
        else:
            section = doc.add_section(WD_SECTION.NEW_PAGE)

        section.page_width  = Pt(self.width)
        section.page_height = Pt(self.height)

        # set page margin
        left,right,top,bottom = self.margin
        section.left_margin = Pt(left)
        section.right_margin = Pt(right)
        section.top_margin = Pt(top)
        section.bottom_margin = Pt(bottom)

        # add paragraph or table according to parsed block
        self.blocks.make_page(doc)


    @debug_plot('Source Text Blocks')
    def plot(self):
        '''Plot initial blocks. It's generally called once Layout is initialized.'''
        return self.blocks

    
    # ----------------------------------------------------
    # initialize layout methods
    # ----------------------------------------------------
    def __source_from_page(self, page):
        '''Source data extracted from page by `PyMuPDF`.'''
        # source blocks
        # NOTE: all these coordinates are relative to un-rotated page
        # https://pymupdf.readthedocs.io/en/latest/page.html#modifying-pages
        raw_layout = page.getText('rawdict')

        # page size: though 'width', 'height' are contained in `raw_dict`, 
        # they are based on un-rotated page. So, update page width/height 
        # to right direction in case page is rotated
        *_, w, h = page.rect # always reflecting page rotation
        raw_layout.update({ 'width' : w, 'height': h })
        self.width, self.height = w, h

        # pre-processing for layout blocks and shapes based on parent page
        self.__preprocess_images(page, raw_layout)
        self.__preprocess_shapes(page, raw_layout)
        
        # BBox is a base class processing coordinates, so set rotation matrix globally
        BBox.set_rotation_matrix(page.rotationMatrix)

        return raw_layout


    def __preprocess_images(self, page, raw):
        ''' Adjust image blocks. Image block extracted by `page.getText('rawdict')` doesn't contain alpha channel data,
            so it needs to get page images by `page.getImageList()` and then recover them. However, `Page.getImageList()` 
            contains each image only once, while `page.getText('rawdict')` generates image blocks for every image location,
            whether or not there are any duplicates. See PyMuPDF doc:
            https://pymupdf.readthedocs.io/en/latest/textpage.html#dictionary-structure-of-extractdict-and-extractrawdict
            
            So, a compromise:
            - get image contents with `page.getImageList()` -> ensure correct images
            - get image location with `page.getText('rawdict')` -> ensure correct locations
        '''
        # recover image blocks
        recovered_images = ImagesExtractor.extract_images(page, self.settings['clip_image_res_ratio'])

        # group original image blocks by image contents
        image_blocks_group = defaultdict(list)
        for block in raw['blocks']:
            if block['type'] != 1: continue
            block['type'] = -1 # "delete" it temporally
            image_blocks_group[hash(block['image'])].append(block)
        
        def same_images(img, img_list):
            bbox = list(map(round, img['bbox']))
            for _img in img_list:
                if list(map(round, _img['bbox']))==bbox: return True
            return False

        # An example to show complicated things here:
        # - images in `page.getImageList`: [a, b, c]
        # - images in `page.getText`     : [a1, a2, b, d]
        # (1) a -> a1, a2: an image in page maps to multi-instances in raw dict
        # (2) c: an image in page may not exist in raw dict -> so, add it
        # (3) d: an image in raw dict may not exist in page -> so, delete it
        for image in recovered_images:
            for k, image_blocks in image_blocks_group.items():
                if not same_images(image, image_blocks): continue
                for image_block in image_blocks:
                    image_block['type'] = 1 # add it back
                    image_block['image'] = image['image']
                break

            # an image outside the page is not counted in page.getText(), so let's add it here
            else:
                raw['blocks'].append(image)


    @debug_plot('Source Paths')
    def __preprocess_shapes(self, page, raw):
        ''' Identify iso-oriented paths and convert vector graphic paths to pixmap.'''
        # extract paths ed by `page.getDrawings()`
        raw_paths = page.getDrawings()

        # paths to shapes or images
        paths = Paths(parent=self).restore(raw_paths)
        images, shapes = paths.to_images_and_shapes(
            page,
            self.settings['curve_path_ratio'], 
            self.settings['clip_image_res_ratio']
            )
        raw['blocks'].extend(images)
        raw['shapes'] = shapes

        return paths


    # ----------------------------------------------------
    # wraping Blocks and Shapes methods
    # ----------------------------------------------------
    @debug_plot('Cleaned Shapes')
    def clean_up_shapes(self):
        '''Clean up shapes and detect semantic types.'''
        # clean up shapes, e.g. remove negative or duplicated instances
        self.shapes.clean_up(self.settings['max_border_width'], 
                            self.settings['shape_merging_threshold'],
                            self.settings['shape_min_dimension'])

        # detect semantic type based on the positions to text blocks, 
        # e.g. table border v.s. text underline, table shading v.s. text highlight.
        # NOTE:
        # stroke shapes are grouped on connectivity to each other, but in some cases, 
        # the gap between borders and underlines/strikes are very close, which leads
        # to an incorrect table structure. So, it's required to distinguish them in
        # advance, though we needn't to ensure 100% accuracy.
        self.shapes.detect_initial_categories()

        return self.shapes


    @debug_plot('Cleaned Blocks')
    def clean_up_blocks(self):
        '''Clean up blocks and calculate page margin accordingly.'''
        # clean up bad blocks, e.g. overlapping, out of page
        self.blocks.clean_up(self.settings['float_image_ignorable_gap'],
                        self.settings['line_overlap_threshold'],
                        self.settings['line_merging_threshold'])

        return self.blocks


    @debug_plot('Lattice Table Structure')
    def parse_lattice_tables(self):
        '''Parse table structure based on explicit stroke shapes.'''
        return self._tables_constructor.lattice_tables(
                self.settings['connected_border_tolerance'],
                self.settings['min_border_clearance'],
                self.settings['max_border_width'],
                self.settings['float_layout_tolerance'],
                self.settings['line_overlap_threshold'],
                self.settings['line_merging_threshold'],
                self.settings['line_separate_threshold']
            )


    @debug_plot('Stream Table Structure')
    def parse_stream_tables(self):
        '''Parse table structure based on layout of blocks.'''
        return self._tables_constructor.stream_tables(
                self.settings['min_border_clearance'],
                self.settings['max_border_width'],
                self.settings['float_layout_tolerance'],
                self.settings['line_overlap_threshold'],
                self.settings['line_merging_threshold'],
                self.settings['line_separate_threshold']
            )


    @debug_plot('Final Layout')
    def parse_text_format(self):
        '''Parse text format in both page and table context.'''
        text_shapes = list(self.shapes.text_underlines_strikes) + list(self.shapes.text_highlights)
        self.blocks.parse_text_format(text_shapes)
        return self.blocks
 

    def parse_spacing(self):
        ''' Calculate external and internal vertical space for paragraph blocks under page context 
            or table context. It'll used as paragraph spacing and line spacing when creating paragraph.
        '''
        self.blocks.parse_spacing(
            self.settings['line_separate_threshold'],
            self.settings['lines_left_aligned_threshold'],
            self.settings['lines_right_aligned_threshold'],
            self.settings['lines_center_aligned_threshold'])

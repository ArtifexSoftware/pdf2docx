'''Table Cell object.'''

from docx.shared import Pt
from ..common.Element import Element
from ..layout.Layout import Layout
from ..common import docx


class Cell(Layout):
    '''Cell object.'''
    def __init__(self, raw:dict=None):
        raw = raw or {}
        super().__init__()
        self.restore(raw) # restore blocks and shapes

        # more cell properties
        self.bg_color     = raw.get('bg_color', None) # type: int
        self.border_color = raw.get('border_color', (0,0,0,0)) # type: tuple [int]
        self.border_width = raw.get('border_width', (0,0,0,0)) # type: tuple [float]
        self.merged_cells = raw.get('merged_cells', (1,1)) # type: tuple [int]


    @property
    def text(self):
        '''Text contained in this cell.'''
        if not self: return None
        # NOTE: sub-table may exists in
        return '\n'.join([block.text if block.is_text_block else '<NEST TABLE>'
                                 for block in self.blocks])


    @property
    def working_bbox(self):
        '''Inner bbox with border excluded.'''
        x0, y0, x1, y1 = self.bbox
        w_top, w_right, w_bottom, w_left = self.border_width
        bbox = (x0+w_left/2.0, y0+w_top/2.0, x1-w_right/2.0, y1-w_bottom/2.0)
        return Element().update_bbox(bbox).bbox # convert to fitz.Rect


    def store(self):
        if not bool(self): return None
        res = super().store()
        res.update({
            'bg_color': self.bg_color,
            'border_color': self.border_color,
            'border_width': self.border_width,
            'merged_cells': self.merged_cells
        })
        return res


    def plot(self, page):
        '''Plot cell and its sub-layout.'''
        super().plot(page)
        self.blocks.plot(page)


    def make_docx(self, table, indexes):
        '''Set cell style and assign contents.

        Args:
            table (Table): ``python-docx`` table instance.
            indexes (tuple): Row and column indexes, ``(i, j)``.
        '''
        # set cell style, e.g. border, shading, cell width
        self._set_style(table, indexes)

        # ignore merged cells
        if not bool(self):  return

        # merge cells
        n_row, n_col = self.merged_cells
        i, j = indexes
        docx_cell = table.cell(i, j)
        if n_row*n_col!=1:
            _cell = table.cell(i+n_row-1, j+n_col-1)
            docx_cell.merge(_cell)

        # ---------------------
        # cell width (cell height is set by row height)
        # ---------------------
        # experience: width of merged cells may change if not setting width for merged cells
        x0, y0, x1, y1 = self.bbox
        docx_cell.width = Pt(x1-x0)

        # insert contents
        # NOTE: there exists an empty paragraph already in each cell, which should be deleted
        # first to avoid unexpected layout. `docx_cell._element.clear_content()` works here.
        # But, docx requires at least one paragraph in each cell, otherwise resulting in a
        # repair error.
        if self.blocks:
            docx_cell._element.clear_content()
            self.blocks.make_docx(docx_cell)


    def _set_style(self, table, indexes):
        '''Set ``python-docx`` cell style, e.g. border, shading, width, row height,
        based on cell block parsed from PDF.

        Args:
            table (Table): ``python-docx`` table object.
            indexes (tuple): ``(i, j)`` index of current cell in table.
        '''
        i, j = indexes
        docx_cell = table.cell(i, j)
        n_row, n_col = self.merged_cells

        # ---------------------
        # border style
        # ---------------------
        # NOTE: border width is specified in eighths of a point, with a minimum value of
        # two (1/4 of a point) and a maximum value of 96 (twelve points)
        keys = ('top', 'end', 'bottom', 'start')
        kwargs = {}
        for k, w, c in zip(keys, self.border_width, self.border_color):
            # skip if width=0 -> will not show in docx
            if not w: continue

            hex_c = f'#{hex(c)[2:].zfill(6)}'
            kwargs[k] = {
                'sz': 8*w, 'val': 'single', 'color': hex_c.upper()
            }

        # merged cells are assumed to have same borders with the main cell
        for m in range(i, i+n_row):
            for n in range(j, j+n_col):
                docx.set_cell_border(table.cell(m, n), **kwargs)

        # ---------------------
        # cell bg-color
        # ---------------------
        if self.bg_color is not None:
            docx.set_cell_shading(docx_cell, self.bg_color)

        # ---------------------
        # clear cell margin
        # ---------------------
        # NOTE: the start position of a table is based on text in cell, rather than
        # left border of table. They're almost aligned if left-margin of cell is zero.
        docx.set_cell_margins(docx_cell, start=0, end=0)

        # set vertical direction if contained text blocks are in vertical direction
        if self.blocks.is_vertical_text:
            docx.set_vertical_cell_direction(docx_cell)
